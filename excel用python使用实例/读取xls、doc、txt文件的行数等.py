#! /usr/lib/python3
# -*- coding: utf-8 -*-
import os,sys

def xls(file):
    import xlrd
    data = xlrd.open_workbook(file)
    table = data.sheets()[0]#第一个工作表
    nrows = table.nrows #行数
    #ncols = table.ncols #列数
    return nrows

def xlsx(file):
    from openpyxl import load_workbook
    wb=load_workbook(file)
    ws=wb.worksheets[0]
    return len(ws.rows)

def docx(file):
    import docx
    #http://www.cnblogs.com/wrajj/p/4914102.html
    '''1、Docment对象表示整个文档；2、Docment包含了Paragraph对象的列表，Paragraph对象用来表示文档中的段落；
    3、一个Paragraph对象包含Run对象的列表，用下面这个图说明Run到底是神马东西。
    Word里面的文本不只是包含了字符串，还有字号、字体、颜色等等属性，都包含在style中。
    一个Run对象就是style相同的一段文本，新建一个Run就有新的style。'''
    doc = docx.Document(file)
    return len(doc.paragraphs)

def txt(file):
    try:
        with open(file)as f:
            return len(f.readlines())
    except UnicodeDecodeError:
        with open(file,'rb')as f:
            return len(f.readlines())

def main():
    file=sys.argv[1]
    print(file)
    #获取文件扩展名
    ex_name=os.path.splitext(file)[-1]
    if ex_name=='.xls':
        print(xls(file))
    elif ex_name=='xlsx':
        print(xlsx(file))
    elif ex_name=='docx':
        print(docx(file))
    else:
        print(txt(file))


if __name__ == "__main__":
    main()

